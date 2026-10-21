"""Project Nura: local-first production orders, planning, release, and execution."""
from __future__ import annotations
import argparse
from contextlib import contextmanager
import io
import json
import os
from pathlib import Path
import secrets
import sqlite3
from datetime import date, datetime, timedelta, timezone
from flask import Flask, jsonify, render_template, request, session, send_file

ROOT = Path(__file__).resolve().parent
STAGES = ('cutting', 'sewing', 'packing')
PLAN_FIELDS = {'line', 'uom', 'size_system', 'scope', 'bom_ref', 'tech_ref', 'qc_ref', 'approver', 'capacity', 'material_ready', 'spec_verified', 'dates_verified', 'material_date', 'cutting_start', 'cutting_end', 'sewing_start', 'sewing_end', 'packing_date', 'delivery_date', 'notes'}
DATES = ('material_date', 'cutting_start', 'cutting_end', 'sewing_start', 'sewing_end', 'packing_date', 'delivery_date')
SCHEMA = '''
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS orders (
 id INTEGER PRIMARY KEY, number TEXT UNIQUE NOT NULL, po TEXT NOT NULL,
 build TEXT NOT NULL, model TEXT NOT NULL, article TEXT NOT NULL, color TEXT NOT NULL,
 buyer TEXT NOT NULL, destination TEXT NOT NULL, gender TEXT NOT NULL, quantity INTEGER NOT NULL CHECK(quantity>0),
 source_packing TEXT, source_delivery TEXT, source_file TEXT, source_row INTEGER
);
CREATE TABLE IF NOT EXISTS sizes(order_id INTEGER REFERENCES orders(id), label TEXT NOT NULL, qty INTEGER NOT NULL CHECK(qty>=0), PRIMARY KEY(order_id,label));
CREATE TABLE IF NOT EXISTS plans(
 order_id INTEGER PRIMARY KEY REFERENCES orders(id), line TEXT NOT NULL DEFAULT '', uom TEXT NOT NULL DEFAULT '',
 size_system TEXT NOT NULL DEFAULT '', scope TEXT NOT NULL DEFAULT '', bom_ref TEXT NOT NULL DEFAULT '',
 tech_ref TEXT NOT NULL DEFAULT '', qc_ref TEXT NOT NULL DEFAULT '', approver TEXT NOT NULL DEFAULT '',
 capacity INTEGER NOT NULL DEFAULT 0, material_ready INTEGER NOT NULL DEFAULT 0, spec_verified INTEGER NOT NULL DEFAULT 0,
 dates_verified INTEGER NOT NULL DEFAULT 0, material_date TEXT NOT NULL DEFAULT '', cutting_start TEXT NOT NULL DEFAULT '',
 cutting_end TEXT NOT NULL DEFAULT '', sewing_start TEXT NOT NULL DEFAULT '', sewing_end TEXT NOT NULL DEFAULT '',
 packing_date TEXT NOT NULL DEFAULT '', delivery_date TEXT NOT NULL DEFAULT '', notes TEXT NOT NULL DEFAULT '',
 version INTEGER NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS work_orders(order_id INTEGER PRIMARY KEY REFERENCES orders(id), status TEXT NOT NULL DEFAULT 'draft', released_at TEXT, released_by TEXT);
CREATE TABLE IF NOT EXISTS events(
 id INTEGER PRIMARY KEY, request_id TEXT UNIQUE NOT NULL, order_id INTEGER NOT NULL REFERENCES orders(id),
 stage TEXT NOT NULL CHECK(stage IN ('cutting','sewing','packing','delivery')),
 kind TEXT NOT NULL CHECK(kind IN ('production','rework','delivery')), event_date TEXT NOT NULL,
 batch TEXT NOT NULL, operator TEXT NOT NULL, input_qty INTEGER NOT NULL CHECK(input_qty>0), good INTEGER NOT NULL CHECK(good>=0),
 rejected INTEGER NOT NULL DEFAULT 0 CHECK(rejected>=0), rework INTEGER NOT NULL DEFAULT 0 CHECK(rework>=0), notes TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, order_id INTEGER REFERENCES orders(id), action TEXT NOT NULL, details TEXT NOT NULL, created_at TEXT NOT NULL);
'''

def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')
def days(start, end):
    a, b = date.fromisoformat(start), date.fromisoformat(end)
    return [a+timedelta(days=i) for i in range((b-a).days+1) if (a+timedelta(days=i)).weekday()<5]

def create_app(database=None, seed=True):
    app = Flask(__name__)
    directory = ROOT / 'instance'
    directory.mkdir(exist_ok=True)
    app.config['DATABASE'] = str(database or os.environ.get('NURA_DB_PATH') or directory/'nura.sqlite3')
    key = directory/'session.key'
    if not key.exists():
        try:
            fd=os.open(key, os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
            with os.fdopen(fd,'w') as f:f.write(secrets.token_hex(32))
        except FileExistsError:pass
    app.secret_key = key.read_text().strip()
    app.config['TRUSTED_HOSTS']=['127.0.0.1','localhost']
    app.config.update(MAX_CONTENT_LENGTH=3*1024*1024, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Strict')
    @contextmanager
    def db():
        c=sqlite3.connect(app.config['DATABASE'],timeout=15)
        c.row_factory=sqlite3.Row
        c.execute('PRAGMA foreign_keys=ON')
        try:
            yield c
            c.commit()
        except Exception:
            c.rollback()
            raise
        finally:
            c.close()
    with db() as c:
        c.executescript(SCHEMA)
        c.execute('PRAGMA user_version=1')
    def log(c,oid,action,details):c.execute('INSERT INTO audit(order_id,action,details,created_at) VALUES(?,?,?,?)',(oid,action,json.dumps(details,ensure_ascii=False),now()))
    def insert(c,d):
        fields=('number','po','build','model','article','color','buyer','destination','gender','quantity','source_packing','source_delivery','source_file','source_row')
        cur=c.execute(f"INSERT INTO orders({','.join(fields)}) VALUES({','.join('?' for _ in fields)})",[d.get(k) for k in fields]);oid=cur.lastrowid
        c.executemany('INSERT INTO sizes VALUES(?,?,?)',[(oid,k,v) for k,v in d['sizes'].items()])
        c.execute('INSERT INTO plans(order_id) VALUES(?)',(oid,));c.execute('INSERT INTO work_orders(order_id) VALUES(?)',(oid,))
        return oid
    with db() as c:
        if seed and not c.execute("SELECT 1 FROM settings WHERE key='initial_seed_v1'").fetchone():
            for d in json.loads((ROOT/'data/initial_orders.json').read_text()):
                if not c.execute('SELECT 1 FROM orders WHERE number=?',(d['number'],)).fetchone(): insert(c,d)
            c.execute("INSERT INTO settings(key,value) VALUES('initial_seed_v1',?)",(now(),))

    def require_order(c,oid):
        d=c.execute('SELECT o.*,w.status,w.released_at,w.released_by FROM orders o JOIN work_orders w ON w.order_id=o.id WHERE o.id=?',(oid,)).fetchone()
        if not d:raise ValueError('Order tidak ditemukan.')
        return dict(d)
    def progress(c,oid):
        out={k:dict(input=0,good=0,rejected=0,rework_pending=0,wip=0) for k in (*STAGES,'delivery')}
        for e in c.execute('SELECT * FROM events WHERE order_id=? ORDER BY id',(oid,)):
            s=out[e['stage']]
            if e['kind']=='rework':s['rework_pending']-=e['input_qty']
            else:s['input']+=e['input_qty']
            s['good']+=e['good'];s['rejected']+=e['rejected'];s['rework_pending']+=e['rework']
        for s in out.values():s['wip']=s['input']-s['good']-s['rejected']-s['rework_pending']
        return out
    def readiness(c,d,p):
        missing=[]
        for f,label in [('line','Line produksi internal'),('uom','Satuan kuantitas'),('size_system','Sistem ukuran'),('scope','Lingkup pekerjaan upper'),('bom_ref','Referensi BOM/revisi'),('tech_ref','Referensi tech pack/pattern'),('qc_ref','Referensi QC'),('approver','Pemberi persetujuan')]:
            if not str(p[f]).strip():missing.append(label+' belum diisi.')
        for f,label in [('material_ready','Kesiapan material'),('spec_verified','Verifikasi spesifikasi'),('dates_verified','Verifikasi arti tanggal sumber')]:
            if not p[f]:missing.append(label+' belum dikonfirmasi.')
        for f in DATES:
            if not p[f]:missing.append('Jadwal '+f.replace('_',' ')+' belum diisi.')
        if p['capacity']<=0:missing.append('Kapasitas sewing harian belum diisi.')
        filled=[(f,p[f]) for f in DATES if p[f]]
        if any(filled[i][1]>filled[i+1][1] for i in range(len(filled)-1)):missing.append('Urutan jadwal harus material → cutting → sewing → packing → delivery.')
        load=None
        if p['sewing_start'] and p['sewing_end'] and p['sewing_start']<=p['sewing_end']:
            if (date.fromisoformat(p['sewing_end'])-date.fromisoformat(p['sewing_start'])).days>366:
                missing.append('Rentang sewing maksimal 366 hari; pecah pekerjaan menjadi SPK yang lebih kecil.')
                return dict(ok=False,issues=missing,capacity=None)
            wd=days(p['sewing_start'],p['sewing_end']);available=len(wd)*p['capacity']
            load=dict(workdays=len(wd),available=available,required=d['quantity'])
            if available<d['quantity']:missing.append(f"Kapasitas sewing periode ini {available:,}, di bawah kuantitas {d['quantity']:,}.")
            if wd and p['line'] and p['capacity']>0:
                mine=d['quantity']/len(wd)
                others=c.execute("SELECT o.quantity,p.* FROM plans p JOIN orders o ON o.id=p.order_id JOIN work_orders w ON w.order_id=o.id WHERE p.line=? AND p.order_id<>? AND w.status IN ('released','in_progress','hold')",(p['line'],d['id'])).fetchall()
                for other in others:
                    if other['uom']!=p['uom'] and other['sewing_start']<=p['sewing_end'] and other['sewing_end']>=p['sewing_start']:
                        missing.append('Satuan antarorder pada line yang sama berbeda. Gunakan line/kapasitas yang satuannya konsisten.');break
                for day in wd:
                    used=mine;caps=[p['capacity']]
                    for other in others:
                        if other['sewing_start']<=day.isoformat()<=other['sewing_end']:
                            od=days(other['sewing_start'],other['sewing_end'])
                            if od:used+=other['quantity']/len(od);caps.append(other['capacity'])
                    if used>min(caps)+.0001:
                        missing.append('Beban sewing line bertabrakan dengan SPK aktif pada '+day.isoformat()+'.');break
        return dict(ok=not missing,issues=missing,capacity=load)
    def detail(c,oid):
        d=require_order(c,oid);d['sizes']=[dict(x) for x in c.execute('SELECT label,qty FROM sizes WHERE order_id=? ORDER BY CAST(label AS REAL)',(oid,))]
        d['plan']=dict(c.execute('SELECT * FROM plans WHERE order_id=?',(oid,)).fetchone());d['progress']=progress(c,oid);d['readiness']=readiness(c,d,d['plan'])
        d['events']=[dict(x) for x in c.execute('SELECT * FROM events WHERE order_id=? ORDER BY id DESC',(oid,))]
        d['audit']=[dict(x) for x in c.execute('SELECT * FROM audit WHERE order_id=? ORDER BY id DESC LIMIT 40',(oid,))]
        return d
    @app.before_request
    def csrf():
        if request.method in ('POST','PUT','DELETE','PATCH'):
            if not secrets.compare_digest(session.get('csrf',''),request.headers.get('X-CSRF-Token','')) or not session.get('csrf'):
                return jsonify(error='Sesi tidak valid. Muat ulang halaman.'),403
    @app.after_request
    def headers(resp):
        resp.headers['X-Content-Type-Options']='nosniff';resp.headers['X-Frame-Options']='DENY';resp.headers['Referrer-Policy']='same-origin'
        resp.headers['Cache-Control']='no-store'
        resp.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        return resp
    @app.errorhandler(ValueError)
    def invalid(e):return jsonify(error=str(e)),400
    @app.errorhandler(413)
    def large(e):return jsonify(error='File melebihi batas 3 MB.'),413
    @app.errorhandler(sqlite3.IntegrityError)
    def constraint(e):return jsonify(error='Data duplikat atau tidak valid. Muat ulang sebelum mencoba lagi.'),409
    @app.get('/')
    def index():
        session.setdefault('csrf',secrets.token_hex(24))
        return render_template('index.html',csrf=session['csrf'])
    @app.get('/api/orders')
    def orders():
        with db() as c:
            rows=[]
            for r in c.execute('SELECT id FROM orders ORDER BY id'):
                d=detail(c,r['id']);d.pop('audit');d.pop('events');rows.append(d)
        return jsonify(rows)
    @app.get('/api/orders/<int:oid>')
    def get_order(oid):
        with db() as c:return jsonify(detail(c,oid))
    @app.put('/api/orders/<int:oid>/plan')
    def plan(oid):
        payload=request.get_json() or {};version=payload.pop('version',None)
        if set(payload)-PLAN_FIELDS:raise ValueError('Kolom planning tidak dikenal.')
        with db() as c:
            c.execute('BEGIN IMMEDIATE');d=require_order(c,oid)
            if d['status']!='draft':raise ValueError('Planning dikunci setelah release. Gunakan HOLD untuk menghentikan pekerjaan; revisi memerlukan SPK baru.')
            existing=dict(c.execute('SELECT * FROM plans WHERE order_id=?',(oid,)).fetchone())
            if version!=existing['version']:return jsonify(error='Data telah berubah. Muat ulang order sebelum menyimpan.'),409
            for f,value in payload.items():
                if f in ('material_ready','spec_verified','dates_verified'):
                    if type(value) not in (bool,int) or value not in (0,1):raise ValueError('Nilai checklist tidak valid.')
                    payload[f]=int(value)
                elif f=='capacity':
                    if type(value)!=int or not 0<=value<=1000000:raise ValueError('Kapasitas harus bilangan bulat 0–1.000.000.')
                else:
                    if not isinstance(value,str) or len(value)>2000:raise ValueError('Teks tidak valid atau terlalu panjang.')
                    payload[f]=value.strip()
                    if f=='line':payload[f]=payload[f].upper()
                    if f in DATES and value:
                        try:
                            parsed=date.fromisoformat(value)
                            if not 2000<=parsed.year<=2100:raise ValueError()
                            payload[f]=parsed.isoformat()
                        except ValueError:raise ValueError('Tanggal harus valid, antara tahun 2000–2100.')
            if payload:c.execute('UPDATE plans SET '+','.join(f'{k}=?' for k in payload)+',version=version+1 WHERE order_id=?',[*payload.values(),oid])
            log(c,oid,'plan_updated',payload)
            return jsonify(detail(c,oid))
    def validate_order(payload):
        fields=('number','po','build','model','article','color','buyer','destination','gender')
        out={}
        for f in fields:
            value=payload.get(f)
            if not isinstance(value,str) or not value.strip() or len(value)>300:raise ValueError('Identitas order wajib diisi: '+f)
            out[f]=value.strip()
        qty=payload.get('quantity')
        if type(qty)!=int or not 0<qty<=1000000:raise ValueError('QTY order harus bilangan bulat 1–1.000.000.')
        sizes=payload.get('sizes')
        if not isinstance(sizes,dict) or not sizes or len(sizes)>40:raise ValueError('Distribusi ukuran wajib diisi.')
        normalized={}
        for label,n in sizes.items():
            try:
                num=float(label)
                if not 0<num<=100 or not (num*2).is_integer():raise ValueError()
                canonical=str(num).removesuffix('.0')
            except Exception:raise ValueError('Label ukuran harus angka atau half size yang valid.')
            if canonical in normalized or type(n)!=int or not 0<=n<=1000000:raise ValueError('Ukuran duplikat atau kuantitas tidak valid.')
            normalized[canonical]=n
        if sum(normalized.values())!=qty:raise ValueError('Total ukuran harus sama dengan QTY order.')
        out.update(quantity=qty,sizes=normalized)
        return out
    @app.post('/api/orders')
    def create_order():
        d=validate_order(request.get_json() or {})
        d.update(source_file='Input manual',source_row=None,source_packing=None,source_delivery=None)
        with db() as c:
            c.execute('BEGIN IMMEDIATE');oid=insert(c,d);log(c,oid,'order_created',d)
            return jsonify(detail(c,oid)),201
    @app.put('/api/orders/<int:oid>')
    def edit_order(oid):
        payload=request.get_json() or {};d=validate_order(payload)
        with db() as c:
            c.execute('BEGIN IMMEDIATE');old=detail(c,oid)
            if old['status']!='draft':raise ValueError('Order dikunci setelah release. Buat SPK baru untuk perubahan yang disetujui.')
            if payload.get('version')!=old['plan']['version']:return jsonify(error='Data berubah. Muat ulang sebelum mengedit.'),409
            fields=[f for f in d if f!='sizes']
            c.execute('UPDATE orders SET '+','.join(f+'=?' for f in fields)+' WHERE id=?',[*(d[f] for f in fields),oid])
            c.execute('DELETE FROM sizes WHERE order_id=?',(oid,))
            c.executemany('INSERT INTO sizes VALUES(?,?,?)',[(oid,k,v) for k,v in d['sizes'].items()])
            c.execute('UPDATE plans SET version=version+1, material_ready=0,spec_verified=0,dates_verified=0 WHERE order_id=?',(oid,))
            log(c,oid,'order_updated',dict(before={**{k:old.get(k) for k in fields},'sizes':old['sizes']},after=d,reason='Checklist direset agar perubahan diverifikasi ulang.'))
            return jsonify(detail(c,oid))
    @app.post('/api/orders/<int:oid>/status')
    def status(oid):
        target=(request.get_json() or {}).get('status')
        with db() as c:
            c.execute('BEGIN IMMEDIATE');d=detail(c,oid);current=d['status']
            if target=='released' and current=='draft':
                if not d['readiness']['ok']:return jsonify(error='SPK belum siap dirilis.',issues=d['readiness']['issues']),422
                c.execute('UPDATE work_orders SET released_at=?,released_by=? WHERE order_id=?',(now(),d['plan']['approver'],oid))
            elif target=='hold' and current in ('released','in_progress'):pass
            elif target=='resume' and current=='hold':target='in_progress' if d['events'] else 'released'
            elif target=='closed' and current in ('released','in_progress'):
                if d['progress']['delivery']['good']!=d['quantity']:raise ValueError('Seluruh kuantitas harus diserahkan sebelum SPK ditutup.')
            else:raise ValueError('Perubahan status tidak diperbolehkan.')
            c.execute('UPDATE work_orders SET status=? WHERE order_id=?',(target,oid));log(c,oid,'status_changed',dict(before=current,after=target))
            return jsonify(detail(c,oid))
    @app.post('/api/orders/<int:oid>/events')
    def events(oid):
        p=request.get_json() or {}
        with db() as c:
            c.execute('BEGIN IMMEDIATE');d=detail(c,oid)
            rid=p.get('request_id','')
            if not isinstance(rid,str) or not 8<=len(rid)<=100:raise ValueError('ID transaksi tidak valid.')
            duplicate=c.execute('SELECT order_id FROM events WHERE request_id=?',(rid,)).fetchone()
            if duplicate:
                if duplicate['order_id']!=oid:raise ValueError('ID transaksi sudah digunakan order lain.')
                return jsonify(detail(c,oid))
            if d['status'] not in ('released','in_progress'):raise ValueError('Aktual hanya dapat dicatat pada SPK aktif yang telah released.')
            stage=p.get('stage');kind=p.get('kind')
            if stage not in (*STAGES,'delivery') or kind not in ('production','rework','delivery'):raise ValueError('Proses/jenis transaksi tidak valid.')
            if (stage=='delivery')!=(kind=='delivery'):raise ValueError('Delivery harus menggunakan jenis delivery.')
            for f in ('input_qty','good','rejected','rework'):
                if type(p.get(f))!=int or not 0<=p[f]<=1000000:raise ValueError('Jumlah harus bilangan bulat nonnegatif.')
            if not p['input_qty'] or p['good']+p['rejected']+p['rework']!=p['input_qty']:raise ValueError('Qty masuk harus sama dengan baik + reject + rework. Catat hanya lot yang selesai diperiksa.')
            for f in ('batch','operator','event_date'):
                if not isinstance(p.get(f),str) or not p[f].strip() or len(p[f])>150:raise ValueError('Batch, PIC dan tanggal wajib diisi.')
            try:date.fromisoformat(p['event_date'])
            except ValueError:raise ValueError('Tanggal tidak valid.')
            if len(str(p.get('notes','')))>2000:raise ValueError('Catatan terlalu panjang.')
            pr=d['progress'];s=pr[stage]
            if kind=='rework':
                if p['input_qty']>s['rework_pending']:raise ValueError('Jumlah melebihi rework yang belum diselesaikan.')
            elif stage=='delivery':
                if p['rejected'] or p['rework'] or p['good']!=p['input_qty']:raise ValueError('Delivery mencatat jumlah diterima, tanpa reject/rework.')
                if p['input_qty']>pr['packing']['good']-s['input']:raise ValueError('Delivery melebihi hasil packing baik yang belum dikirim.')
            else:
                maximum=d['quantity'] if stage=='cutting' else pr[STAGES[STAGES.index(stage)-1]]['good']
                if s['input']+p['input_qty']>maximum:raise ValueError('Jumlah melebihi order atau hasil baik proses sebelumnya. Produksi pengganti memerlukan order tambahan yang disetujui.')
            if d['events'] and p['event_date']<max(e['event_date'] for e in d['events']):raise ValueError('Tanggal transaksi tidak boleh lebih awal dari transaksi terakhir order ini.')
            c.execute('INSERT INTO events(request_id,order_id,stage,kind,event_date,batch,operator,input_qty,good,rejected,rework,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(rid,oid,stage,kind,p['event_date'],p['batch'].strip(),p['operator'].strip(),p['input_qty'],p['good'],p['rejected'],p['rework'],str(p.get('notes','')),now()))
            c.execute("UPDATE work_orders SET status='in_progress' WHERE order_id=?",(oid,));log(c,oid,'actual_recorded',p)
            return jsonify(detail(c,oid)),201
    def parse_workbook(file):
        from openpyxl import load_workbook
        import zipfile
        raw=file.read()
        try:
            z=zipfile.ZipFile(io.BytesIO(raw))
            if sum(i.file_size for i in z.infolist())>25*1024*1024:raise ValueError('Isi workbook terlalu besar.')
            w=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
        except ValueError:raise
        except Exception:raise ValueError('File bukan workbook XLSX yang dapat dibaca.')
        name='PROD PLAN ORDER BIG - SEP 2026'
        if name not in w.sheetnames:raise ValueError('Gunakan template DATA ORDER DAN SPK NURA: sheet '+name+'.')
        s=w[name]
        if s.max_row>5000 or s.max_column>100:raise ValueError('Dimensi workbook melebihi batas impor.')
        for cell,label in [('B4','Work Order/SPK'),('D4','Nomor PO'),('K4','QTY Order'),('AN5','Packing'),('AO5','Delivery')]:
            if s[cell].value!=label:raise ValueError('Header template tidak cocok pada '+cell+'.')
        labels=[]
        for col in range(12,35):
            try:labels.append(str(float(s.cell(5,col).value)).removesuffix('.0'))
            except Exception:raise ValueError('Header ukuran tidak valid.')
        if len(set(labels))!=23:raise ValueError('Label ukuran duplikat.')
        rows=[];seen=set()
        for rowno,r in enumerate(s.iter_rows(min_row=6,max_col=41),6):
            if all(c.value is None for c in r):continue
            vals=[c.value for c in r]
            if any(vals[i] is None for i in [1,2,3,4,5,6,7,8,9,10]):raise ValueError(f'Baris {rowno}: identitas wajib belum lengkap.')
            qty=vals[10]
            if type(qty) not in (int,float) or qty<=0 or int(qty)!=qty:raise ValueError(f'Baris {rowno}: QTY tidak valid.')
            ss={}
            for i,label in enumerate(labels,11):
                value=vals[i] if vals[i] is not None else 0
                if type(value) not in (int,float) or value<0 or int(value)!=value:raise ValueError(f'Baris {rowno}: kuantitas ukuran tidak valid.')
                ss[label]=int(value)
            if sum(ss.values())!=qty:raise ValueError(f'Baris {rowno}: jumlah ukuran tidak cocok dengan QTY.')
            number=str(vals[1]).strip()
            if not number or number in seen:raise ValueError(f'Baris {rowno}: nomor SPK duplikat/kosong.')
            seen.add(number)
            row=dict(zip(('number','build','po','model','article','color','buyer','destination','gender'),[str(x).strip() for x in vals[1:10]]))
            row.update(quantity=int(qty),sizes=ss,source_file=Path(file.filename or 'import.xlsx').name,source_row=rowno)
            for field,ix in [('source_packing',39),('source_delivery',40)]:
                value=vals[ix]
                if value is not None and not isinstance(value,(datetime,date)):raise ValueError(f'Baris {rowno}: tanggal sumber tidak valid.')
                row[field]=value.date().isoformat() if isinstance(value,datetime) else value.isoformat() if value else None
            rows.append(row)
        w.close()
        if not rows:raise ValueError('Tidak ada order dalam file.')
        return rows
    @app.post('/api/import')
    def import_orders():
        if 'file' not in request.files:raise ValueError('Pilih file XLSX terlebih dahulu.')
        rows=parse_workbook(request.files['file']);apply=request.form.get('apply')=='1'
        with db() as c:
            c.execute('BEGIN IMMEDIATE');existing={r['number']:r for r in c.execute('SELECT * FROM orders')};new=[];same=0;conflicts=[]
            for d in rows:
                old=existing.get(d['number'])
                if not old:new.append(d);continue
                ss={r['label']:r['qty'] for r in c.execute('SELECT * FROM sizes WHERE order_id=?',(old['id'],))}
                if any(str(old[f])!=str(d[f]) for f in ('po','build','model','article','color','buyer','destination','gender','quantity','source_packing','source_delivery')) or ss!=d['sizes']:conflicts.append(d['number'])
                else:same+=1
            if conflicts:return jsonify(error='Ada nomor SPK sama dengan data berbeda. Impor dibatalkan; data lama tidak ditimpa.',conflicts=conflicts),409
            if apply:
                for d in new:oid=insert(c,d);log(c,oid,'imported',dict(file=d['source_file'],row=d['source_row']))
            return jsonify(total=len(rows),new=len(new),unchanged=same,quantity=sum(d['quantity'] for d in rows),applied=apply)
    @app.get('/api/backup')
    def backup():
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'nura-backup.sqlite3'
            with db() as src,sqlite3.connect(path) as dst:src.backup(dst)
            data=path.read_bytes()
        return send_file(io.BytesIO(data),as_attachment=True,download_name='nura-backup-'+date.today().isoformat()+'.sqlite3',mimetype='application/vnd.sqlite3')
    return app

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--port',type=int,default=5000);args=parser.parse_args()
    create_app().run(host='127.0.0.1',port=args.port,debug=False)
