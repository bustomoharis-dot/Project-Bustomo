import io
import json
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import create_app

class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name)/'test.sqlite3'
        self.app=create_app(self.path);self.app.testing=True;self.c=self.app.test_client();self.c.get('/')
        with self.c.session_transaction() as sess:self.headers={'X-CSRF-Token':sess['csrf']}
    def tearDown(self):self.tmp.cleanup()
    def get(self,oid=1):return self.c.get(f'/api/orders/{oid}').json
    def post(self,url,data):return self.c.post(url,json=data,headers=self.headers)
    def plan(self,oid=1,**changes):
        d=self.get(oid);p=dict(version=d['plan']['version'],line='SEW-01',uom='pair',size_system='UK Men',scope='Upper lengkap',bom_ref='BOM-01 Rev A',tech_ref='TP-01 Rev A',qc_ref='QC-01 Rev A',approver='Supervisor',capacity=155,material_ready=True,spec_verified=True,dates_verified=True,material_date='2026-09-21',cutting_start='2026-09-21',cutting_end='2026-09-21',sewing_start='2026-09-21',sewing_end='2026-09-21',packing_date='2026-09-22',delivery_date='2026-09-23',notes='Test')
        p.update(changes);r=self.c.put(f'/api/orders/{oid}/plan',json=p,headers=self.headers);self.assertEqual(r.status_code,200,r.json);return r.json
    def release(self,oid=1):
        self.plan(oid);r=self.post(f'/api/orders/{oid}/status',{'status':'released'});self.assertEqual(r.status_code,200,r.json)
    def event(self,stage,qty,good=None,kind='production',rid=None,rework=0,rejected=0):
        return self.post('/api/orders/1/events',dict(request_id=rid or f'request-{stage}-{kind}-{qty}-{good}',stage=stage,kind=kind,event_date='2026-09-23',batch='BATCH-01',operator='Admin',input_qty=qty,good=qty if good is None else good,rejected=rejected,rework=rework,notes=''))
    def workbook(self,conflict=False):
        w=Workbook();s=w.active;s.title='PROD PLAN ORDER BIG - SEP 2026'
        for cell,value in [('B4','Work Order/SPK'),('D4','Nomor PO'),('K4','QTY Order'),('AN5','Packing'),('AO5','Delivery')]:s[cell]=value
        seed=json.loads((Path(__file__).resolve().parents[1]/'data/initial_orders.json').read_text())
        labels=list(seed[0]['sizes'])
        for i,l in enumerate(labels,12):s.cell(5,i,float(l))
        from datetime import datetime
        for r,d in enumerate(seed,6):
            for col,key in enumerate(['number','build','po','model','article','color','buyer','destination','gender','quantity'],2):s.cell(r,col,d[key])
            if conflict and r==6:s.cell(r,5,'CHANGED MODEL')
            for col,l in enumerate(labels,12):s.cell(r,col,d['sizes'][l])
            for col,key in [(40,'source_packing'),(41,'source_delivery')]:s.cell(r,col,datetime.fromisoformat(d[key]))
        out=io.BytesIO();w.save(out);out.seek(0);return out
    def test_seed_integrity_and_persistence(self):
        rows=self.c.get('/api/orders').json;self.assertEqual(len(rows),59);self.assertEqual(sum(o['quantity'] for o in rows),34741)
        self.assertTrue(all(sum(x['qty'] for x in o['sizes'])==o['quantity'] for o in rows));self.assertTrue(all(o['status']=='draft' and not o['plan']['packing_date'] for o in rows))
        create_app(self.path);self.assertEqual(len(self.c.get('/api/orders').json),59)
    def test_seed_not_reinserted_after_source_number_edit(self):
        d=self.get();payload={k:d[k] for k in ['number','po','build','model','article','color','buyer','destination','gender','quantity']}
        payload.update(number='RENAMED-SOURCE',version=1,sizes={s['label']:s['qty'] for s in d['sizes']})
        r=self.c.put('/api/orders/1',json=payload,headers=self.headers);self.assertEqual(r.status_code,200)
        create_app(self.path);rows=self.c.get('/api/orders').json
        self.assertEqual(len(rows),59);self.assertEqual(rows[0]['number'],'RENAMED-SOURCE')
    def test_csrf_is_required(self):
        self.assertEqual(self.c.post('/api/orders/1/status',json={'status':'released'}).status_code,403)
    def test_incomplete_release_blocked(self):
        r=self.post('/api/orders/1/status',{'status':'released'});self.assertEqual(r.status_code,422);self.assertGreater(len(r.json['issues']),10)
        self.assertEqual(self.event('cutting',1).status_code,400)
    def test_optimistic_lock_and_schedule_validation(self):
        self.plan();r=self.c.put('/api/orders/1/plan',json={'version':1,'line':'OTHER'},headers=self.headers);self.assertEqual(r.status_code,409)
        self.plan(sewing_start='2026-09-25',sewing_end='2026-09-21');self.assertFalse(self.get()['readiness']['ok'])
    def test_capacity_and_shared_line_collision(self):
        self.plan(capacity=100);self.assertFalse(self.get()['readiness']['ok']);self.release()
        self.plan(2);r=self.post('/api/orders/2/status',{'status':'released'});self.assertEqual(r.status_code,422);self.assertTrue(any('bertabrakan' in x for x in r.json['issues']))
    def test_stage_flow_limits_and_closed_order(self):
        self.release();self.assertEqual(self.event('sewing',1).status_code,400);self.assertEqual(self.event('cutting',156).status_code,400)
        for s in ['cutting','sewing','packing']:self.assertEqual(self.event(s,155).status_code,201)
        self.assertEqual(self.event('delivery',156,kind='delivery').status_code,400)
        self.assertEqual(self.event('delivery',155,kind='delivery').status_code,201)
        self.assertEqual(self.post('/api/orders/1/status',{'status':'closed'}).status_code,200)
        self.assertEqual(self.event('cutting',1).status_code,400)
    def test_rework_and_duplicate_request_not_double_counted(self):
        self.release();self.assertEqual(self.event('cutting',155,150,rework=5,rid='unique-initial').status_code,201)
        self.assertEqual(self.event('cutting',155,150,rework=5,rid='unique-initial').status_code,200)
        self.assertEqual(self.get()['progress']['cutting']['input'],155)
        self.assertEqual(self.event('cutting',6,kind='rework').status_code,400)
        self.assertEqual(self.event('cutting',5,kind='rework').status_code,201)
        p=self.get()['progress']['cutting'];self.assertEqual((p['input'],p['good'],p['rework_pending']),(155,155,0))
    def test_hold_and_lock_after_release(self):
        self.release();self.assertEqual(self.post('/api/orders/1/status',{'status':'hold'}).status_code,200)
        self.assertEqual(self.event('cutting',1).status_code,400)
        self.assertEqual(self.c.put('/api/orders/1/plan',json={'version':2,'line':'X'},headers=self.headers).status_code,400)
        self.assertEqual(self.post('/api/orders/1/status',{'status':'resume'}).status_code,200)
    def test_cannot_close_undelivered(self):
        self.release();self.assertEqual(self.post('/api/orders/1/status',{'status':'closed'}).status_code,400)
    def test_import_idempotence_and_conflict_atomicity(self):
        for apply in ['0','1']:
            r=self.c.post('/api/import',data={'file':(self.workbook(),'orders.xlsx'),'apply':apply},headers=self.headers);self.assertEqual(r.status_code,200,r.json);self.assertEqual((r.json['new'],r.json['unchanged']),(0,59))
        r=self.c.post('/api/import',data={'file':(self.workbook(True),'orders.xlsx'),'apply':'1'},headers=self.headers);self.assertEqual(r.status_code,409)
        self.assertEqual(self.get()['model'],'FRONTCOURT LOW')
    def test_manual_order_validation_and_draft_edit(self):
        d=dict(number='TEST-01',po='000123',build='CUSTOMER-BUILD',model='Test',article='ART',color='WHITE',buyer='Buyer',destination='ID',gender='M',quantity=10,sizes={'8':10})
        r=self.post('/api/orders',d);self.assertEqual(r.status_code,201,r.json);self.assertEqual(r.json['po'],'000123');oid=r.json['id']
        self.plan(oid,line='OTHER');d.update(version=2,quantity=11,sizes={'8':11})
        r=self.c.put(f'/api/orders/{oid}',json=d,headers=self.headers);self.assertEqual(r.status_code,200,r.json);self.assertEqual(r.json['plan']['spec_verified'],0)
        d['quantity']=12;self.assertEqual(self.post('/api/orders',d).status_code,400)
    def test_backup_is_readable_sqlite(self):
        r=self.c.get('/api/backup');self.assertEqual(r.status_code,200)
        path=Path(self.tmp.name)/'backup.sqlite3';path.write_bytes(r.data)
        with sqlite3.connect(path) as c:self.assertEqual(c.execute('SELECT SUM(quantity) FROM orders').fetchone()[0],34741)

if __name__=='__main__':unittest.main()
